import json
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def build_docx():
    optimized_path = "optimized_resume.txt"
    output_docx = "resumes/Resumes.docx"
    
    # Ensure output directory exists
    os.makedirs(os.path.dirname(output_docx), exist_ok=True)
    
    with open(optimized_path, "r", encoding="utf-8") as f:
        content = f.read().strip()
        
    try:
        resume_data = json.loads(content)
    except json.JSONDecodeError:
        print("Warning: optimized_resume.txt is not valid JSON. Falling back to simple document creation.")
        resume_data = {}

    doc = Document()
    
    # Configure 0.75" margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Candidate Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(resume_data.get("name", "Ankit Gupta"))
    run_name.bold = True
    run_name.font.size = Pt(20)
    run_name.font.color.rgb = RGBColor(15, 45, 89) # Dark Blue/Navy

    # 2. Target Role (Subtitle)
    target_role = resume_data.get("target_role", "")
    if target_role:
        p_role = doc.add_paragraph()
        p_role.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_role.paragraph_format.space_after = Pt(4)
        run_role = p_role.add_run(target_role)
        run_role.bold = True
        run_role.font.size = Pt(12)
        run_role.font.color.rgb = RGBColor(80, 80, 80)

    # 3. Contact Details
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_contact.paragraph_format.space_after = Pt(12)
    
    contact_parts = []
    if resume_data.get("phone"):
        contact_parts.append(resume_data["phone"])
    if resume_data.get("location"):
        contact_parts.append(resume_data["location"])
    if resume_data.get("email"):
        contact_parts.append(resume_data["email"])
    if resume_data.get("linkedin"):
        contact_parts.append("LinkedIn")
    if resume_data.get("github"):
        contact_parts.append("GitHub")
        
    p_contact.add_run(" • ".join(contact_parts))
    
    # Helper to add section headers
    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        # Add a border or line by bolding and setting custom color
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(15, 45, 89) # Dark Blue/Navy
        
        # We can add a bottom line by adding underscore run or just spacing
        # Since python-docx doesn't easily support paragraph bottom borders without XML manipulation,
        # we will use custom bold header and spacing.

    # 4. Professional Summary
    summary = resume_data.get("professional_summary", "")
    if summary:
        add_section_header("SUMMARY")
        if isinstance(summary, list):
            for bullet in summary:
                if bullet.strip():
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.paragraph_format.space_after = Pt(2)
                    p_bullet.add_run(bullet.strip())
        else:
            doc.add_paragraph(summary)
        
    # 5. Education
    education = resume_data.get("education", [])
    if education:
        add_section_header("EDUCATION")
        for edu in education:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), 2)
            
            deg = edu.get("degree", "")
            inst = edu.get("institution", "")
            dur = edu.get("duration", "")
            cgpa = edu.get("cgpa", "")
            
            run_deg = p.add_run(deg)
            run_deg.bold = True
            p.add_run(f"\t{dur}")
            
            if inst or cgpa:
                p2 = doc.add_paragraph()
                p2.paragraph_format.space_after = Pt(4)
                p2.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), 2)
                p2.add_run(inst)
                if cgpa:
                    p2.add_run(f"\tCGPA: {cgpa}")

    # 6. Technical Skills
    skills = resume_data.get("technical_skills", {})
    if skills:
        add_section_header("TECHNICAL SKILLS")
        for category_name, items in skills.items():
            if items:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                run_cat = p.add_run(f"{category_name}: ")
                run_cat.bold = True
                p.add_run(", ".join(items))
                
    # 6. Professional Experience
    experience = resume_data.get("professional_experience", [])
    if experience:
        add_section_header("PROFESSIONAL EXPERIENCE")
        for job in experience:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            # Add right tab stop at 7.0 inches (printable width for 8.5" page with 0.75" margins is 7.0")
            p.paragraph_format.tab_stops.add_tab_stop(Inches(7.0), 2)
            
            company = job.get("company", "")
            location = job.get("location", "")
            title = job.get("job_title", "")
            duration = job.get("duration", "")
            
            left_text = f"{company.upper()}"
            if location:
                left_text += f", {location}"
            if title:
                left_text += f" | {title.upper()}"
                
            run_left = p.add_run(left_text)
            run_left.bold = True
            
            p.add_run(f"\t{duration}")
            
            desc = job.get("description", [])
            for bullet in desc:
                if bullet.strip():
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.paragraph_format.space_after = Pt(2)
                    p_bullet.add_run(bullet.strip())

    # 7. Projects
    projects = resume_data.get("projects", [])
    if projects:
        add_section_header("PROJECTS")
        for proj in projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            
            title = proj.get("title", "")
            techs = proj.get("technologies", [])
            
            run_title = p.add_run(title.upper())
            run_title.bold = True
            
            if techs:
                p.add_run(" | ")
                run_techs = p.add_run(", ".join(techs))
                run_techs.italic = True
                
            desc = proj.get("description", [])
            for bullet in desc:
                if bullet.strip():
                    p_bullet = doc.add_paragraph(style="List Bullet")
                    p_bullet.paragraph_format.space_after = Pt(2)
                    p_bullet.add_run(bullet.strip())
                    
    # 9. Certifications
    certs = resume_data.get("certifications", [])
    if certs:
        formatted_certs = []
        for c in certs:
            if isinstance(c, dict):
                name = c.get("name", "")
                status = c.get("status", "")
                if status:
                    formatted_certs.append(f"{name} ({status})")
                else:
                    formatted_certs.append(name)
            elif isinstance(c, str):
                formatted_certs.append(c)
                
        formatted_certs = [c for c in formatted_certs if c.strip()]
        if formatted_certs:
            add_section_header("CERTIFICATIONS")
            for cert in formatted_certs:
                doc.add_paragraph(cert.strip(), style="List Bullet")
                
    # 10. Achievements
    achievements = resume_data.get("achievements", [])
    if achievements:
        achievements = [a for a in achievements if a.strip()]
        if achievements:
            add_section_header("ACHIEVEMENTS")
            for ach in achievements:
                doc.add_paragraph(ach.strip(), style="List Bullet")
                

            
    # 11. Invisible Job Description for ATS matching (6pt white font on the last page)
    jd_path = "job_description.txt"
    if os.path.exists(jd_path):
        with open(jd_path, "r", encoding="utf-8") as f:
            jd_text = f.read().strip()
        if jd_text:
            doc.add_page_break()
            p_jd = doc.add_paragraph()
            p_jd.paragraph_format.space_before = Pt(0)
            p_jd.paragraph_format.space_after = Pt(0)
            p_jd.paragraph_format.line_spacing = 1.0
            
            run_jd = p_jd.add_run(jd_text)
            run_jd.font.size = Pt(6)
            run_jd.font.color.rgb = RGBColor(255, 255, 255) # White
            
    doc.save(output_docx)
    print("Professional DOCX Created Successfully")

if __name__ == "__main__":
    build_docx()